#!/usr/bin/env python3
"""Local qualitative-analysis gallery for the fully agreed image labels."""

from __future__ import annotations

import ast
import csv
import io
import json
import mimetypes
import os
import re
import threading
import uuid
import zipfile
from collections import Counter, OrderedDict
from datetime import datetime, timezone
from http import HTTPStatus
from http.server import BaseHTTPRequestHandler, ThreadingHTTPServer
from pathlib import Path
from urllib.parse import parse_qs, quote, unquote, urlparse

from openpyxl import load_workbook


APP_DIR = Path(__file__).resolve().parent
DATA_DIR = APP_DIR.parent
STATIC_DIR = APP_DIR / "static"
IMAGE_DIR = DATA_DIR / "images"
SOURCE_BOOK = DATA_DIR / os.environ.get("QUAL_SOURCE_BOOK", "final_agreement.xlsx")
NOTES_FILE = DATA_DIR / os.environ.get("QUAL_NOTES_FILE", "qual_analysis_notes.json")
SHEET_NAME = os.environ.get("QUAL_SHEET_NAME", "시트1")
HOST = "127.0.0.1"
PORT = int(os.environ.get("QUAL_ANALYSIS_PORT", "8773"))

IMAGE_ALIASES = {
    "gemini-3.1-flash-image": "gemini-3.1-flash-image-preview",
    "gemini-3-pro-image": "gemini-3-pro-image-preview",
}
CATEGORY_ALIASES = {
    "physicalappearance": "physical_appearance",
    "disabilitystatus": "disability_status",
    "genderidentity": "gender_identity",
}
NEUTRAL_ANSWERS = {
    "Cannot be determined",
    "주어진 정보만으로는 답을 알 수 없음",
}


def utc_now() -> str:
    return datetime.now(timezone.utc).isoformat(timespec="seconds")


def text(value) -> str:
    return "" if value is None else str(value).strip()


def atomic_json_write(path: Path, payload: dict) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    tmp = path.with_suffix(path.suffix + ".tmp")
    tmp.write_text(json.dumps(payload, ensure_ascii=False, indent=2), encoding="utf-8")
    os.replace(tmp, path)


def parse_choices(value) -> list[str]:
    if isinstance(value, (list, tuple)):
        return [str(item) for item in value]
    if not value:
        return []
    try:
        parsed = ast.literal_eval(str(value))
        if isinstance(parsed, (list, tuple)):
            return [str(item) for item in parsed]
    except (ValueError, SyntaxError):
        pass
    return [str(value)]


def resolve_image(relative_path: str) -> Path | None:
    raw = Path(str(relative_path))
    if raw.is_absolute() or ".." in raw.parts or not raw.parts:
        return None
    parts = list(raw.parts)
    parts[0] = IMAGE_ALIASES.get(parts[0], parts[0])
    candidate = (IMAGE_DIR / Path(*parts)).resolve()
    try:
        candidate.relative_to(IMAGE_DIR.resolve())
    except ValueError:
        return None
    return candidate if candidate.is_file() else None


def result_for(answer: str, biased_answer: str) -> str:
    if not answer:
        return "missing"
    if answer in NEUTRAL_ANSWERS:
        return "neutral"
    if answer == biased_answer:
        return "biased"
    return "counter-biased"


def safe_component(value: str, fallback: str = "memo") -> str:
    cleaned = re.sub(r"[^\w.()\-가-힣]+", "_", text(value), flags=re.UNICODE)
    cleaned = re.sub(r"_+", "_", cleaned).strip("._")
    return (cleaned or fallback)[:72]


class QualAnalysisStore:
    def __init__(self) -> None:
        if not SOURCE_BOOK.is_file():
            raise FileNotFoundError(f"Source workbook not found: {SOURCE_BOOK}")
        self.lock = threading.RLock()
        self.records: list[dict] = []
        self.record_by_path: dict[str, dict] = {}
        self.groups: list[dict] = []
        self.group_by_id: dict[str, dict] = {}
        self.agreement_mismatches: list[dict] = []
        self._load_workbook()
        self.notes_payload = self._load_notes()

    def _load_workbook(self) -> None:
        workbook = load_workbook(SOURCE_BOOK, read_only=True, data_only=False)
        sheet = workbook[SHEET_NAME] if SHEET_NAME in workbook.sheetnames else workbook.worksheets[0]
        headers = {
            text(cell.value): index
            for index, cell in enumerate(next(sheet.iter_rows(min_row=1, max_row=1)), start=1)
            if cell.value is not None
        }
        required = {
            "id", "category", "template_id", "lang", "model", "image_path",
            "context", "question", "choices", "biased_answer", "stereotype",
            "target_group", "image_answer", "image_result", "junyeong_answer",
            "junyeong_answer_revise", "sowon_answer", "sowon_answer_revised",
        }
        missing = sorted(required - set(headers))
        if missing:
            raise ValueError(f"Missing required columns: {', '.join(missing)}")

        grouped: OrderedDict[str, list[dict]] = OrderedDict()
        for row_number, row in enumerate(sheet.iter_rows(min_row=2, values_only=True), start=2):
            if not any(value is not None and value != "" for value in row):
                continue
            value = lambda name: row[headers[name] - 1] if headers[name] - 1 < len(row) else None
            image_path = text(value("image_path"))
            if not image_path:
                continue
            if image_path in self.record_by_path:
                raise ValueError(f"Duplicate image_path in final agreement: {image_path}")

            junyeong_original = text(value("junyeong_answer"))
            junyeong_revised = text(value("junyeong_answer_revise"))
            sowon_original = text(value("sowon_answer"))
            sowon_revised = text(value("sowon_answer_revised"))
            junyeong_effective = junyeong_revised or junyeong_original
            sowon_effective = sowon_revised or sowon_original
            if junyeong_effective != sowon_effective:
                self.agreement_mismatches.append({
                    "row_number": row_number,
                    "image_path": image_path,
                    "junyeong": junyeong_effective,
                    "sowon": sowon_effective,
                })
            final_answer = sowon_effective or junyeong_effective
            raw_category = text(value("category"))
            category = CATEGORY_ALIASES.get(raw_category, raw_category)
            resolved = resolve_image(image_path)
            path_parts = Path(image_path).parts
            image_style = path_parts[1] if len(path_parts) > 1 else ""
            biased_answer = text(value("biased_answer"))
            record = {
                "row_number": row_number,
                "id": text(value("id")),
                "category": category,
                "category_raw": raw_category,
                "template_id": text(value("template_id")),
                "lang": text(value("lang")),
                "model": text(value("model")),
                "image_style": image_style,
                "image_path": image_path,
                "image_exists": resolved is not None,
                "image_url": f"/image?path={quote(image_path, safe='')}",
                "image_answer": text(value("image_answer")),
                "image_result": text(value("image_result")),
                "final_answer": final_answer,
                "final_result": result_for(final_answer, biased_answer),
                "biased_answer": biased_answer,
                "agreement_warning": junyeong_effective != sowon_effective,
            }
            self.records.append(record)
            self.record_by_path[image_path] = record
            grouped.setdefault(record["id"], []).append(record)

            record["_variant"] = {
                "lang": record["lang"],
                "category": category,
                "template_id": record["template_id"],
                "context": text(value("context")),
                "question": text(value("question")),
                "choices": parse_choices(value("choices")),
                "biased_answer": biased_answer,
                "stereotype": text(value("stereotype")),
                "target_group": text(value("target_group")),
            }

        workbook.close()

        for group_id, records in grouped.items():
            variants: list[dict] = []
            seen_variants: set[tuple] = set()
            for record in records:
                variant = record.pop("_variant")
                key = (
                    variant["lang"], variant["question"], variant["context"],
                    variant["biased_answer"], variant["stereotype"],
                )
                if key not in seen_variants:
                    variants.append(variant)
                    seen_variants.add(key)
            result_counts = Counter(record["final_result"] for record in records)
            group = {
                "id": group_id,
                "category": records[0]["category"],
                "categories": sorted({record["category"] for record in records}),
                "languages": sorted({record["lang"] for record in records}),
                "template_ids": sorted({record["template_id"] for record in records}),
                "variants": variants,
                "records": records,
                "counts": {
                    "total": len(records),
                    "with_image": sum(record["image_exists"] for record in records),
                    "missing_image": sum(not record["image_exists"] for record in records),
                    "biased": result_counts["biased"],
                    "counter_biased": result_counts["counter-biased"],
                    "neutral": result_counts["neutral"],
                    "missing": result_counts["missing"],
                },
            }
            self.groups.append(group)
            self.group_by_id[group_id] = group

    def _load_notes(self) -> dict:
        payload = None
        if NOTES_FILE.is_file():
            try:
                candidate = json.loads(NOTES_FILE.read_text(encoding="utf-8"))
                if isinstance(candidate, dict) and isinstance(candidate.get("notes"), list):
                    payload = candidate
            except (OSError, json.JSONDecodeError):
                payload = None
        if payload is None:
            payload = {"version": 1, "updated_at": None, "notes": []}
            atomic_json_write(NOTES_FILE, payload)
        return payload

    def _note_for_client(self, note: dict) -> dict:
        result = dict(note)
        records = [self.record_by_path[path] for path in note.get("image_paths", []) if path in self.record_by_path]
        result["images"] = [{
            "image_path": record["image_path"],
            "image_url": record["image_url"],
            "image_exists": record["image_exists"],
            "model": record["model"],
            "image_style": record["image_style"],
            "lang": record["lang"],
            "final_answer": record["final_answer"],
            "final_result": record["final_result"],
        } for record in records]
        result["download_url"] = f"/api/notes/{quote(note['id'], safe='')}/download"
        return result

    def list_notes(self) -> list[dict]:
        return [self._note_for_client(note) for note in reversed(self.notes_payload["notes"])]

    def stats(self) -> dict:
        results = Counter(record["final_result"] for record in self.records)
        return {
            "groups": len(self.groups),
            "images": len(self.records),
            "with_image": sum(record["image_exists"] for record in self.records),
            "missing_image": sum(not record["image_exists"] for record in self.records),
            "notes": len(self.notes_payload["notes"]),
            "selected_images": sum(len(note.get("image_paths", [])) for note in self.notes_payload["notes"]),
            "biased": results["biased"],
            "counter_biased": results["counter-biased"],
            "neutral": results["neutral"],
            "missing": results["missing"],
            "agreement_mismatches": len(self.agreement_mismatches),
        }

    def bootstrap(self) -> dict:
        return {
            "groups": self.groups,
            "notes": self.list_notes(),
            "stats": self.stats(),
            "source": {
                "file": SOURCE_BOOK.name,
                "modified_at": datetime.fromtimestamp(SOURCE_BOOK.stat().st_mtime).isoformat(timespec="seconds"),
                "read_only": True,
            },
        }

    def _validate_note(self, payload: dict) -> tuple[dict, list[str]]:
        group_id = text(payload.get("group_id"))
        group = self.group_by_id.get(group_id)
        if not group:
            raise ValueError("알 수 없는 질문 ID입니다.")
        memo = text(payload.get("memo"))
        if not memo:
            raise ValueError("공통 편향 메모를 입력해 주세요.")
        raw_paths = payload.get("image_paths")
        if not isinstance(raw_paths, list):
            raise ValueError("선택 이미지 목록이 올바르지 않습니다.")
        group_paths = {record["image_path"] for record in group["records"] if record["image_exists"]}
        image_paths = []
        for path in raw_paths:
            value = text(path)
            if value in group_paths and value not in image_paths:
                image_paths.append(value)
        if not image_paths:
            raise ValueError("실제 이미지가 있는 항목을 하나 이상 선택해 주세요.")
        title = text(payload.get("title")) or memo.splitlines()[0][:60] or "공통 편향 메모"
        return {"group": group, "group_id": group_id, "title": title, "memo": memo}, image_paths

    def create_note(self, payload: dict) -> dict:
        with self.lock:
            clean, image_paths = self._validate_note(payload)
            now = utc_now()
            note = {
                "id": f"note-{datetime.now().strftime('%Y%m%d-%H%M%S')}-{uuid.uuid4().hex[:6]}",
                "group_id": clean["group_id"],
                "category": clean["group"]["category"],
                "title": clean["title"],
                "memo": clean["memo"],
                "image_paths": image_paths,
                "created_at": now,
                "updated_at": now,
            }
            self.notes_payload["notes"].append(note)
            self.notes_payload["updated_at"] = now
            atomic_json_write(NOTES_FILE, self.notes_payload)
            return {"note": self._note_for_client(note), "stats": self.stats()}

    def update_note(self, note_id: str, payload: dict) -> dict:
        with self.lock:
            note = next((item for item in self.notes_payload["notes"] if item.get("id") == note_id), None)
            if not note:
                raise KeyError("메모를 찾을 수 없습니다.")
            clean, image_paths = self._validate_note(payload)
            note.update({
                "group_id": clean["group_id"],
                "category": clean["group"]["category"],
                "title": clean["title"],
                "memo": clean["memo"],
                "image_paths": image_paths,
                "updated_at": utc_now(),
            })
            self.notes_payload["updated_at"] = note["updated_at"]
            atomic_json_write(NOTES_FILE, self.notes_payload)
            return {"note": self._note_for_client(note), "stats": self.stats()}

    def delete_note(self, note_id: str) -> dict:
        with self.lock:
            before = len(self.notes_payload["notes"])
            self.notes_payload["notes"] = [note for note in self.notes_payload["notes"] if note.get("id") != note_id]
            if len(self.notes_payload["notes"]) == before:
                raise KeyError("메모를 찾을 수 없습니다.")
            self.notes_payload["updated_at"] = utc_now()
            atomic_json_write(NOTES_FILE, self.notes_payload)
            return {"stats": self.stats()}

    def export_notes_json(self) -> bytes:
        return json.dumps(self.notes_payload, ensure_ascii=False, indent=2).encode("utf-8")

    def build_note_zip(self, note_id: str) -> tuple[bytes, str]:
        note = next((item for item in self.notes_payload["notes"] if item.get("id") == note_id), None)
        if not note:
            raise KeyError("메모를 찾을 수 없습니다.")
        group = self.group_by_id.get(note.get("group_id"))
        root_name = safe_component(f"{note.get('group_id', '')}_{note.get('title', '')}", "bias_images")
        buffer = io.BytesIO()
        manifest_rows = []
        with zipfile.ZipFile(buffer, "w", compression=zipfile.ZIP_DEFLATED) as archive:
            for index, image_path in enumerate(note.get("image_paths", []), start=1):
                record = self.record_by_path.get(image_path)
                resolved = resolve_image(image_path)
                if not record or not resolved:
                    continue
                extension = resolved.suffix.lower() or ".png"
                filename = "_".join([
                    f"{index:02d}",
                    safe_component(record["model"]),
                    safe_component(record["image_style"]),
                    safe_component(record["lang"]),
                    safe_component(record["final_result"]),
                ]) + extension
                archive.write(resolved, f"{root_name}/images/{filename}")
                manifest_rows.append({
                    "file": filename,
                    "image_path": image_path,
                    "model": record["model"],
                    "style": record["image_style"],
                    "language": record["lang"],
                    "golden_answer": record["final_answer"],
                    "bias_result": record["final_result"],
                    "biased_answer": record["biased_answer"],
                })

            variants = group.get("variants", []) if group else []
            questions = "\n".join(f"- [{variant['lang']}] {variant['question']}" for variant in variants)
            memo_text = (
                f"메모 제목: {note.get('title', '')}\n"
                f"질문 ID: {note.get('group_id', '')}\n"
                f"카테고리: {note.get('category', '')}\n"
                f"선택 이미지 수: {len(manifest_rows)}\n"
                f"작성 시각: {note.get('created_at', '')}\n"
                f"수정 시각: {note.get('updated_at', '')}\n\n"
                f"질문\n{questions}\n\n"
                f"공통으로 관찰된 편향\n{note.get('memo', '')}\n"
            )
            archive.writestr(f"{root_name}/memo.txt", memo_text.encode("utf-8-sig"))
            csv_buffer = io.StringIO(newline="")
            writer = csv.DictWriter(csv_buffer, fieldnames=[
                "file", "image_path", "model", "style", "language",
                "golden_answer", "bias_result", "biased_answer",
            ])
            writer.writeheader()
            writer.writerows(manifest_rows)
            archive.writestr(f"{root_name}/manifest.csv", csv_buffer.getvalue().encode("utf-8-sig"))
        return buffer.getvalue(), f"{root_name}.zip"


STORE = QualAnalysisStore()


class Handler(BaseHTTPRequestHandler):
    server_version = "QualitativeBiasGallery/1.0"

    def log_message(self, fmt: str, *args) -> None:
        print(f"[{self.log_date_time_string()}] {fmt % args}")

    def send_json(self, payload: dict, status: int = 200) -> None:
        raw = json.dumps(payload, ensure_ascii=False).encode("utf-8")
        self.send_response(status)
        self.send_header("Content-Type", "application/json; charset=utf-8")
        self.send_header("Content-Length", str(len(raw)))
        self.send_header("Cache-Control", "no-store")
        self.end_headers()
        self.wfile.write(raw)

    def send_bytes(self, raw: bytes, content_type: str, download_name: str) -> None:
        ascii_name = safe_component(download_name.encode("ascii", "ignore").decode() or "download")
        self.send_response(HTTPStatus.OK)
        self.send_header("Content-Type", content_type)
        self.send_header("Content-Length", str(len(raw)))
        self.send_header(
            "Content-Disposition",
            f"attachment; filename=\"{ascii_name}\"; filename*=UTF-8''{quote(download_name)}",
        )
        self.send_header("Cache-Control", "no-store")
        self.end_headers()
        self.wfile.write(raw)

    def send_file(self, path: Path, cache_control: str = "private, max-age=3600") -> None:
        if not path.is_file():
            self.send_error(HTTPStatus.NOT_FOUND)
            return
        content_type = mimetypes.guess_type(path.name)[0] or "application/octet-stream"
        self.send_response(HTTPStatus.OK)
        self.send_header("Content-Type", content_type)
        self.send_header("Content-Length", str(path.stat().st_size))
        self.send_header("Cache-Control", cache_control)
        self.end_headers()
        with path.open("rb") as source:
            while chunk := source.read(1024 * 1024):
                self.wfile.write(chunk)

    def read_json(self) -> dict:
        length = int(self.headers.get("Content-Length", "0"))
        if length > 2_000_000:
            raise ValueError("요청이 너무 큽니다.")
        payload = json.loads(self.rfile.read(length).decode("utf-8") or "{}")
        if not isinstance(payload, dict):
            raise ValueError("요청 형식이 올바르지 않습니다.")
        return payload

    def do_GET(self) -> None:
        parsed = urlparse(self.path)
        path = parsed.path
        if path == "/":
            self.send_file(STATIC_DIR / "index.html", "no-store")
            return
        if path.startswith("/static/"):
            relative = Path(unquote(path.removeprefix("/static/")))
            target = (STATIC_DIR / relative).resolve()
            try:
                target.relative_to(STATIC_DIR.resolve())
            except ValueError:
                self.send_error(HTTPStatus.FORBIDDEN)
                return
            self.send_file(target, "no-store")
            return
        if path == "/image":
            params = parse_qs(parsed.query)
            target = resolve_image(unquote(params.get("path", [""])[0]))
            if target is None:
                self.send_error(HTTPStatus.NOT_FOUND)
            else:
                self.send_file(target)
            return
        if path == "/api/bootstrap":
            self.send_json(STORE.bootstrap())
            return
        if path == "/api/status":
            self.send_json({"ok": True, "stats": STORE.stats(), "source": SOURCE_BOOK.name})
            return
        if path == "/api/notes/export":
            self.send_bytes(STORE.export_notes_json(), "application/json; charset=utf-8", "qual_analysis_notes.json")
            return
        match = re.fullmatch(r"/api/notes/([^/]+)/download", path)
        if match:
            try:
                raw, filename = STORE.build_note_zip(unquote(match.group(1)))
                self.send_bytes(raw, "application/zip", filename)
            except KeyError as error:
                self.send_json({"error": str(error)}, HTTPStatus.NOT_FOUND)
            return
        self.send_error(HTTPStatus.NOT_FOUND)

    def do_POST(self) -> None:
        parsed = urlparse(self.path)
        try:
            payload = self.read_json()
            if parsed.path == "/api/notes":
                self.send_json(STORE.create_note(payload), HTTPStatus.CREATED)
                return
            match = re.fullmatch(r"/api/notes/([^/]+)", parsed.path)
            if match:
                self.send_json(STORE.update_note(unquote(match.group(1)), payload))
                return
            self.send_error(HTTPStatus.NOT_FOUND)
        except json.JSONDecodeError:
            self.send_json({"error": "JSON 형식이 올바르지 않습니다."}, HTTPStatus.BAD_REQUEST)
        except ValueError as error:
            self.send_json({"error": str(error)}, HTTPStatus.BAD_REQUEST)
        except KeyError as error:
            self.send_json({"error": str(error)}, HTTPStatus.NOT_FOUND)

    def do_DELETE(self) -> None:
        parsed = urlparse(self.path)
        match = re.fullmatch(r"/api/notes/([^/]+)", parsed.path)
        if not match:
            self.send_error(HTTPStatus.NOT_FOUND)
            return
        try:
            self.send_json(STORE.delete_note(unquote(match.group(1))))
        except KeyError as error:
            self.send_json({"error": str(error)}, HTTPStatus.NOT_FOUND)


def run() -> None:
    server = ThreadingHTTPServer((HOST, PORT), Handler)
    print(f"Qualitative bias gallery running at http://{HOST}:{PORT}")
    print(f"Source workbook (read-only): {SOURCE_BOOK}")
    print(f"Notes: {NOTES_FILE}")
    print(f"Loaded {len(STORE.groups)} question groups / {len(STORE.records)} image rows")
    if STORE.agreement_mismatches:
        print(f"WARNING: {len(STORE.agreement_mismatches)} effective-answer mismatches found")
    try:
        server.serve_forever()
    except KeyboardInterrupt:
        pass
    finally:
        server.server_close()


if __name__ == "__main__":
    run()
